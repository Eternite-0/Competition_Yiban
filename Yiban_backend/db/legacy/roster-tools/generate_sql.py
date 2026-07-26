import openpyxl
import re

wb = openpyxl.load_workbook(r'D:\易班\副本2024级学生信息表(867).xlsx', data_only=True)
ws = wb.active

# 读取表头
headers = [cell.value for cell in ws[1]]
print(f"表头: {headers}")

# 解析数据
students = []
for row in ws.iter_rows(min_row=2, values_only=True):
    if row[0] is None:  # 跳过空行
        continue

    name = str(row[0]).strip()  # 姓名
    gender = str(row[2]).strip() if row[2] else ''  # 性别
    grade = str(row[3]).strip() if row[3] else '2024'  # 年级
    major = str(row[4]).strip() if row[4] else ''  # 专业名称
    class_name = str(row[5]).strip() if row[5] else ''  # 班级
    student_no = str(row[6]).strip() if row[6] else ''  # 学号

    if student_no and name:
        students.append({
            'student_no': student_no,
            'name': name,
            'college': '计算机与人工智能学院',
            'major': major,
            'class_name': class_name,
            'grade': grade
        })

print(f"解析到 {len(students)} 条学生数据")

# 提取唯一的专业和班级
majors = set()
classes = set()
for s in students:
    majors.add(s['major'])
    classes.add((s['class_name'], s['major'], s['grade']))

print(f"专业数量: {len(majors)}")
print(f"班级数量: {len(classes)}")

# 生成 SQL
sql_lines = []
sql_lines.append("-- =============================================")
sql_lines.append("-- 学生花名册导入脚本")
sql_lines.append(f"-- 生成时间: 2026-05-30")
sql_lines.append(f"-- 学生总数: {len(students)}")
sql_lines.append("-- =============================================")
sql_lines.append("")
sql_lines.append("USE `etsaion`;")
sql_lines.append("")

# 1. 插入专业
sql_lines.append("-- 1. 插入专业数据")
sql_lines.append("INSERT INTO `major` (`name`, `college`, `status`) VALUES")
major_values = []
for major in sorted(majors):
    major_values.append(f"('{major}', '计算机与人工智能学院', 'active')")
sql_lines.append(",\n".join(major_values) + ";")
sql_lines.append("")

# 2. 插入班级
sql_lines.append("-- 2. 插入班级数据（关联专业）")
sql_lines.append("INSERT INTO `class_info` (`name`, `college`, `major_id`, `grade`, `status`) VALUES")
class_values = []
for class_name, major, grade in sorted(classes):
    class_values.append(f"('{class_name}', '计算机与人工智能学院', (SELECT id FROM major WHERE name='{major}'), '{grade}', 'active')")
sql_lines.append(",\n".join(class_values) + ";")
sql_lines.append("")

# 3. 插入花名册
sql_lines.append("-- 3. 插入学生花名册数据")
sql_lines.append("INSERT INTO `student_roster` (`student_no`, `real_name`, `college`, `major_id`, `class_id`, `grade`, `status`) VALUES")
roster_values = []
for s in students:
    roster_values.append(
        f"('{s['student_no']}', '{s['name']}', '计算机与人工智能学院', "
        f"(SELECT id FROM major WHERE name='{s['major']}'), "
        f"(SELECT id FROM class_info WHERE name='{s['class_name']}' AND grade='{s['grade']}'), "
        f"'{s['grade']}', 'pending')"
    )
sql_lines.append(",\n".join(roster_values) + ";")
sql_lines.append("")

# 4. 验证
sql_lines.append("-- 验证导入结果")
sql_lines.append("SELECT 'major' AS table_name, COUNT(*) AS count FROM `major`")
sql_lines.append("UNION ALL")
sql_lines.append("SELECT 'class_info', COUNT(*) FROM `class_info`")
sql_lines.append("UNION ALL")
sql_lines.append("SELECT 'student_roster', COUNT(*) FROM `student_roster`;")

# 写入文件
sql_content = "\n".join(sql_lines)
with open(r'D:\易班\Competition_Yiban\Yiban_backend\db\import-roster-2024.sql', 'w', encoding='utf-8') as f:
    f.write(sql_content)

print(f"\nSQL 脚本已生成: Yiban_backend/db/import-roster-2024.sql")
print(f"包含 {len(majors)} 个专业, {len(classes)} 个班级, {len(students)} 名学生")
