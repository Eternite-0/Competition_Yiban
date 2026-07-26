import openpyxl

wb = openpyxl.load_workbook(r'D:\易班\副本2024级学生信息表(867).xlsx')
ws = wb.active

print(f'行数: {ws.max_row}')
print(f'列数: {ws.max_column}')

print('\n表头:')
headers = [cell.value for cell in ws[1]]
print(headers)

print('\n前10行数据:')
for i, row in enumerate(ws.iter_rows(min_row=2, max_row=11, values_only=True), 2):
    print(f'行{i}: {row}')

print(f'\n总数据行数: {ws.max_row - 1}')
